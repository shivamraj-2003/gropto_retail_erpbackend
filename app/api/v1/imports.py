import io
import uuid

from fastapi import APIRouter, Depends, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import CurrentUser, get_current_user, require_permission, require_store_access, require_warehouse_access
from app.core.database import get_db
from app.models.models import ImportBatch
from app.services import imports as import_service

router = APIRouter(prefix="/imports", tags=["imports"])


TEMPLATES = {
    "products": {
        "permission": "import.product.import",
        "sheet": "Products",
        "file": "product-import-template.xlsx",
        "columns": ["sku", "name", "barcode", "uom", "purchase_price", "selling_price", "mrp", "tax_rate", "hsn_code", "other_barcodes"],
        "examples": [
            ["EXAMPLE-001", "Cadbury Bournvita 500g (example - delete this row)", "8901233012345", "EA", 215, 240, 250, 18, "1806", "8901233099999, 8901233088888"],
            ["EXAMPLE-002", "Tata Salt 1kg (example - delete this row)", "8901000000002", "EA", 18, 22, 25, 5, "2501", ""],
        ],
        "help": [
            ("sku", "Yes", "Your own unique code for the product. This is how the file finds an existing product.", "BOURN-500"),
            ("name", "New products", "The name shown on bills and screens. Leave blank on an existing product to keep its name.", "Cadbury Bournvita 500g"),
            ("barcode", "No", "The main barcode. Must not belong to another product.", "8901233012345"),
            ("uom", "No", "Unit: EA (each), KG, G, L, ML, PKT, BOX, DOZ.", "EA"),
            ("purchase_price", "No", "What you pay the vendor, per unit.", "215"),
            ("selling_price", "New products", "Price the customer pays, including GST. Changing it on an existing product goes for approval (Super Admin: applied at once).", "240"),
            ("mrp", "No", "The printed price. Cannot be lower than the selling price. Empty = same as selling price.", "250"),
            ("tax_rate", "No", "GST percentage: 0, 5, 12, 18 or 28.", "18"),
            ("hsn_code", "No", "HSN code for GST invoices.", "1806"),
            ("other_barcodes", "No", "More barcodes of the same product, separated by commas. Scanning any of them finds the product. Only added, never removed.", "8901233099999, 8901233088888"),
        ],
        "notes": [
            "A blank cell never changes anything: only the cells you fill in are updated.",
            "Delete the two EXAMPLE rows before you upload.",
            "You will see a row-by-row preview first. Nothing is saved until you press Save.",
        ],
    },
    "opening-stock": {
        "permission": "import.opening_stock.import",
        "sheet": "Opening stock",
        "file": "opening-stock-template.xlsx",
        "columns": ["sku", "quantity", "batch_number", "mfg_date", "expiry_date"],
        "examples": [
            ["EXAMPLE-001", 120, "B2410", "2026-09-01", "2027-03-01"],
            ["EXAMPLE-002", 300, "", "", ""],
        ],
        "help": [
            ("sku", "Yes", "The SKU of a product that already exists (add it in Products first).", "BOURN-500"),
            ("quantity", "Yes", "How many units are on the shelf now. Must be more than 0.", "120"),
            ("batch_number", "No", "Batch number, for items that expire.", "B2410"),
            ("mfg_date", "No", "Manufacturing date: YYYY-MM-DD or DD-MM-YYYY.", "2026-09-01"),
            ("expiry_date", "No", "Expiry date. Fill batch and dates for perishables so expiry is tracked.", "2027-03-01"),
        ],
        "notes": [
            "Pick the store or warehouse on the Import Data screen - it is not a column in the file.",
            "Opening stock can be loaded once per product and place. After that use Purchases, Transfers or a stock count.",
            "Delete the two EXAMPLE rows before you upload.",
        ],
    },
}


@router.get("/templates/{kind}")
async def download_template(
    kind: str,
    current: CurrentUser = Depends(get_current_user),
) -> StreamingResponse:
    """A ready-to-fill Excel file: the right columns, two example rows, and a "How to fill" sheet."""
    spec = TEMPLATES.get(kind)
    if spec is None:
        raise HTTPException(status_code=404, detail="Unknown template")
    if not current.has_permission(spec["permission"]):
        raise HTTPException(status_code=403, detail=f"Permission denied: {spec['permission']}")
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = spec["sheet"]
    ws.append(spec["columns"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
    for row in spec["examples"]:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(italic=True, color="9CA3AF")
    for idx, column in enumerate(spec["columns"], start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = max(14, len(column) + 6)
    ws.column_dimensions["B"].width = 42
    ws.freeze_panes = "A2"

    guide = wb.create_sheet("How to fill")
    guide.append(["Column", "Needed?", "What to put", "Example"])
    for cell in guide[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
    for row in spec["help"]:
        guide.append(list(row))
    guide.append([])
    for note in spec["notes"]:
        guide.append([note])
    for col, width in zip("ABCD", (18, 14, 90, 34)):
        guide.column_dimensions[col].width = width
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{spec["file"]}"'},
    )


@router.post("/products/stage")
async def stage_products(
    file: UploadFile,
    db: AsyncSession = Depends(get_db),
    current: CurrentUser = Depends(require_permission("import.product.import")),
) -> dict:
    content = await file.read()
    try:
        batch = await import_service.stage_products_file(db, file_bytes=content, filename=file.filename, uploaded_by=current.user_id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return await import_service.preview_batch(db, batch.id)


@router.get("/{batch_id}/preview")
async def preview(
    batch_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    _current: CurrentUser = Depends(require_permission("import.batch.view")),
) -> dict:
    return await import_service.preview_batch(db, batch_id)


@router.post("/products/{batch_id}/commit")
async def commit_products(
    batch_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current: CurrentUser = Depends(require_permission("import.product.import")),
) -> dict:
    try:
        return await import_service.commit_products_batch(db, batch_id, current)
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc


@router.post("/opening-stock/stage")
async def stage_opening_stock(
    file: UploadFile,
    store_id: uuid.UUID | None = None,
    warehouse_id: uuid.UUID | None = None,
    db: AsyncSession = Depends(get_db),
    current: CurrentUser = Depends(require_permission("import.opening_stock.import")),
) -> dict:
    """Opening stock for one store OR one warehouse (exactly one of the two)."""
    if (store_id is None) == (warehouse_id is None):
        raise HTTPException(status_code=400, detail="Pick either a store or a warehouse to load stock into")
    if store_id is not None:
        require_store_access(store_id, current)
    else:
        require_warehouse_access(warehouse_id, current)
    content = await file.read()
    try:
        batch = await import_service.stage_opening_stock_file(
            db, file_bytes=content, filename=file.filename, store_id=store_id, warehouse_id=warehouse_id, uploaded_by=current.user_id
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return await import_service.preview_batch(db, batch.id)


@router.post("/opening-stock/{batch_id}/commit")
async def commit_opening_stock(
    batch_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current: CurrentUser = Depends(require_permission("import.opening_stock.import")),
) -> dict:
    try:
        return await import_service.commit_opening_stock_batch(db, batch_id, current)
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc


@router.get("/{batch_id}/file")
async def download_original_file(
    batch_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    _current: CurrentUser = Depends(require_permission("import.batch.view")),
) -> StreamingResponse:
    """Every batch keeps its uploaded file (§10) — this is how you get it back."""
    batch = await db.get(ImportBatch, batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail="Batch not found")
    return StreamingResponse(
        io.BytesIO(batch.file_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{batch.file_name}"'},
    )


@router.get("/{batch_id}/result-report")
async def download_result_report(
    batch_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    _current: CurrentUser = Depends(require_permission("import.batch.view")),
) -> StreamingResponse:
    """A downloadable result report showing exactly what was applied, skipped and
    rejected, with row numbers (§10)."""
    batch = await import_service.get_batch_with_rows(db, batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail="Batch not found")

    wb = Workbook()
    ws = wb.active
    ws.title = "Import Result"
    ws.append(["Row", "Action", "Errors", "Parsed Values"])
    for row in batch.rows:
        ws.append([row.line_number, row.computed_action, "; ".join(row.validation_messages), str(row.parsed_values)])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="import_result_{batch_id}.xlsx"'},
    )
